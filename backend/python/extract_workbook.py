"""Extract an xlsx workbook into the JSON contract consumed by TypeScript."""

from __future__ import annotations

import json
import sys
from datetime import date, datetime, time
from io import BytesIO
from typing import Any

from openpyxl import load_workbook


MAX_WORKSHEETS = 100
MAX_ROWS_PER_SHEET = 50_000
MAX_COLUMNS_PER_SHEET = 256
MAX_CELLS_PER_WORKBOOK = 1_000_000


def color_text(color: Any) -> str:
    color_type = getattr(color, "type", None)
    value = getattr(color, color_type, None) if color_type else None
    if value is None:
        return ""
    return str(value).upper()


def cell_text(cell: Any, cached_cell: Any) -> str:
    value = cell.value
    if cell.data_type == "f":
        cached_value = cached_cell.value
        if cached_value is not None:
            value = cached_value
        elif value is not None:
            formula = str(value)
            return formula if formula.startswith("=") else f"={formula}"

    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def extract(payload: bytes, file_name: str) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(payload), data_only=False, read_only=False)
    cached_workbook = load_workbook(BytesIO(payload), data_only=True, read_only=True)

    if len(workbook.worksheets) > MAX_WORKSHEETS:
        raise ValueError(f"Workbook has more than {MAX_WORKSHEETS} worksheets.")

    total_cells = 0
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        row_count = worksheet.max_row or 0
        column_count = worksheet.max_column or 0
        if row_count > MAX_ROWS_PER_SHEET:
            raise ValueError(
                f'Worksheet "{worksheet.title}" exceeds the '
                f"{MAX_ROWS_PER_SHEET:,} row limit."
            )
        if column_count > MAX_COLUMNS_PER_SHEET:
            raise ValueError(
                f'Worksheet "{worksheet.title}" exceeds the '
                f"{MAX_COLUMNS_PER_SHEET} column limit."
            )

        total_cells += row_count * column_count
        if total_cells > MAX_CELLS_PER_WORKBOOK:
            raise ValueError(
                f"Workbook exceeds the {MAX_CELLS_PER_WORKBOOK:,} cell limit."
            )

        cached_sheet = cached_workbook[worksheet.title]
        cached_rows = cached_sheet.iter_rows(
            min_row=1,
            max_row=row_count,
            min_col=1,
            max_col=column_count,
        )
        matrix: list[list[str]] = []
        row_meta: list[dict[str, Any]] = []
        formulas: list[dict[str, Any]] = []

        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=row_count,
                min_col=1,
                max_col=column_count,
            ),
            start=1,
        ):
            cached_row = next(cached_rows)
            values: list[str] = []
            fill_colors: list[str] = []
            for cell, cached_cell in zip(row, cached_row):
                values.append(cell_text(cell, cached_cell))
                fill_colors.append(color_text(cell.fill.fgColor))
                if cell.data_type == "f" and cell.value is not None:
                    formula = str(cell.value)
                    formulas.append(
                        {
                            "cell": cell.coordinate,
                            "formula": formula if formula.startswith("=") else f"={formula}",
                        }
                    )

            matrix.append(values)
            row_meta.append(
                {
                    "excelRowNumber": row_number,
                    "hidden": bool(worksheet.row_dimensions[row_number].hidden),
                    "fillColors": fill_colors,
                }
            )

        hidden_columns = [
            column_name
            for column_name, dimension in worksheet.column_dimensions.items()
            if dimension.hidden
        ]
        sheets.append(
            {
                "sheetName": worksheet.title,
                "matrix": matrix,
                "rowMeta": row_meta,
                "mergedRanges": [str(cell_range) for cell_range in worksheet.merged_cells.ranges],
                "hiddenColumns": hidden_columns,
                "formulas": formulas,
            }
        )

    workbook.close()
    cached_workbook.close()
    return {"fileName": file_name, "sheets": sheets}


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    file_name = sys.argv[1] if len(sys.argv) > 1 else "workbook.xlsx"
    try:
        result = extract(sys.stdin.buffer.read(), file_name)
        json.dump(result, sys.stdout, ensure_ascii=False, separators=(",", ":"))
        return 0
    except Exception as error:  # The TypeScript adapter turns stderr into a safe error.
        print(f"{type(error).__name__}: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
